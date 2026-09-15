"""Parity against actual formulas and all usable saved FX-sheet P&L cells.

FX current caches contain Excel errors, so they cannot provide numerical parity.
The 26 numeric saved historical futures values DO verify the formula including K.
"""
from pathlib import Path

import openpyxl
import pytest

from engine.pnl.pnl import workbook_fx_pnl, workbook_valuation_date
from engine.pnl.pnl import ltd_per_trade
from data.ingest.schema import connect

WORKBOOK = Path(__file__).resolve().parents[1] / "data/raw/HA-portfolio vJean.xlsx"
pytestmark = pytest.mark.skipif(not WORKBOOK.exists(), reason="reference workbook not present")


def test_database_pipeline_matches_saved_historical_futures():
    book = openpyxl.load_workbook(WORKBOOK, data_only=True)
    sheet = book['All FX trades']
    previous = sheet['N1'].value.date().isoformat()
    previous2 = sheet['O1'].value.date().isoformat()
    maturity = sheet['P1'].value.date().isoformat()
    conn = connect()
    expected = {}
    for r in range(2, sheet.max_row + 1):
        name, c, fill = sheet.cell(r,2).value, sheet.cell(r,3).value, sheet.cell(r,5).value
        td = sheet.cell(r,1).value
        if name not in ('ESU6 Index','ESZ6 Index') or td.date().isoformat() > previous2:
            continue
        expiry = sheet.cell(r,4).value.date().isoformat()
        conn.execute('INSERT OR IGNORE INTO instruments VALUES (?,?,?,?,?,?,?,?)',
                     (name,'FUTURE','ES','USD',50,0,name,expiry))
        trade_id = str(r)
        quantity = c/(50*fill)
        conn.execute('INSERT INTO trades VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                     (trade_id,'XLSX',name,'FUTURE',trade_id,td.date().isoformat(),quantity,fill,'a','b','s','t','cached parity fixture',''))
        conn.execute('INSERT INTO trade_legs VALUES (?,?,?,?,?,?,?,?,?)',
                     (trade_id,1,'NOTIONAL','USD',c,td.date().isoformat(),expiry,fill,0))
        for day,column in [(previous,6),(previous2,10)]:
            conn.execute('INSERT OR IGNORE INTO marks VALUES (?,?,?,?,?,?,?)',
                         (day,name,expiry,'FUTURE_PX',sheet.cell(r,column).value,'WORKBOOK_REFERENCE',day+'T15:00:00-04:00'))
        expected[trade_id] = (sheet.cell(r,8).value,sheet.cell(r,11).value)
    conn.commit()
    h = ltd_per_trade(conn,previous,'WORKBOOK_REFERENCE',valuation_date=maturity).set_index('trade_id')
    k = ltd_per_trade(conn,previous2,'WORKBOOK_REFERENCE',valuation_date=maturity,denominator_as_of=previous).set_index('trade_id')
    assert len(expected) >= 8
    for trade_id,(saved_h,saved_k) in expected.items():
        assert h.loc[trade_id,'pnl_usd'] == pytest.approx(saved_h, abs=1e-8)
        assert k.loc[trade_id,'pnl_usd'] == pytest.approx(saved_k, abs=1e-8)
    conn.close()
    book.close()


def test_all_numeric_cached_historical_values_match_literal_formula():
    formula = openpyxl.load_workbook(WORKBOOK, data_only=False)["All FX trades"]
    cached = openpyxl.load_workbook(WORKBOOK, data_only=True)["All FX trades"]
    checked = []
    for row in range(2, cached.max_row + 1):
        name = cached.cell(row, 2).value
        for output, valuation in [(8, 6), (11, 10)]:
            expected = cached.cell(row, output).value
            if not isinstance(expected, (float, int)):
                continue
            assert isinstance(formula.cell(row, output).value, str)
            actual = workbook_fx_pnl(name, cached.cell(row, 3).value,
                cached.cell(row, 5).value, cached.cell(row, valuation).value,
                cached.cell(row, 6).value)
            assert actual == pytest.approx(expected, abs=1e-8), cached.cell(row, output).coordinate
            checked.append(cached.cell(row, output).coordinate)
    assert len(checked) == 26, "Do not silently lose cached numerical reference coverage"


def test_actual_fx_formulas_and_saved_shared_date():
    formulas = openpyxl.load_workbook(WORKBOOK, data_only=False)["All FX trades"]
    values = openpyxl.load_workbook(WORKBOOK, data_only=True)["All FX trades"]
    assert formulas["P1"].value == "=WORKDAY(M1, 5)"
    assert workbook_valuation_date(values["M1"].value.date().isoformat()) == values["P1"].value.date().isoformat()
    count = 0
    for row in range(2, formulas.max_row + 1):
        pair = values.cell(row, 2).value
        if not isinstance(pair, str) or len(pair) != 6:
            continue
        assert formulas.cell(row, 9).value == f'=$C{row}*(G{row}-$E{row})/ IF(RIGHT($B{row},3)="USD", $E{row},G{row})'
        assert formulas.cell(row, 11).value == f'=$C{row}*(J{row}-$E{row})/ IF(RIGHT($B{row},3)="USD", $E{row},F{row})'
        assert not isinstance(values.cell(row, 9).value, (int, float)), "New numeric FX cache requires parity coverage"
        count += 1
    assert count > 300
