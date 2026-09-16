"""Futures fills from the workbook's 'All FX trades' sheet.

The BNP snapshot carries only one netted position per futures contract with an average
cost (see data/ingest/bnp.py module docstring), never per-fill price or date, so futures
trades / trade_legs must come from the xlsx blotter instead. FX rows of the same sheet
are still NOT imported (CLAUDE.md: "FX rows of the workbook are still NOT imported").

A futures row is any 'All FX trades' row whose pair column ends in ' Index'. Column
'Quantity' (canonically C) holds a saved formula of the shape '=<n>*E<row>*50' or
'=<n>*50*E<row>' (the two orders both occur in the reference workbook) recovering the
contract count `n` from the fill column and the fixed multiplier 50. If the workbook was
saved without the formula (a plain cached number only -- openpyxl always sees the formula
string when one is present, data_only load is not needed here) contracts are instead
recovered as `quantity / (50 * fill)`.

Format tolerance (2026-09-16, in response to uploads failing on cosmetic formatting):
this parser is deliberately lenient about anything cosmetic -- sheet name case/whitespace,
header whitespace/case, header column order, blank padding rows/columns, and date
representation (Excel date, ISO string, US mm/dd/yyyy string, or Excel serial number). It
is NOT lenient about the data itself: a value that cannot be parsed as a date/number, or a
pair that doesn't match the documented futures symbol shape, is never guessed at -- the row
is skipped and recorded as a RowIssue with a plain-English message, per CLAUDE.md's "never
silently coerce or guess incorrect values" rule and the reconciliation tolerances that
depend on exact fills. The rest of the file still loads.
"""
from __future__ import annotations

import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

import openpyxl
import openpyxl.utils.datetime as xl_datetime

from data.ingest import schema
from data.ingest.bnp import third_friday

SHEET = "All FX trades"
FUTURES_FORMULA_RE = re.compile(r"^=(-?\d+)\*E\d+\*50$|^=(-?\d+)\*50\*E\d+$")
INDEX_PAIR_RE = re.compile(r"^([A-Z]+)([FGHJKMNQUVXZ])(\d) Index$")
MONTH_CODES = "FGHJKMNQUVXZ"

# Canonical A-E layout per CLAUDE.md ("All FX trades" A Date, B pair, C Quantity, D tenor,
# E fill) used only as a fallback when a column's header cell is blank / unrecognised --
# the reference workbook itself leaves the pair column (B) header blank.
_DEFAULT_COLUMNS = {"date": 1, "pair": 2, "quantity": 3, "settle": 4, "fill": 5}

_HEADER_ALIASES = {
    "date": {"date", "trade date"},
    "pair": {"pair", "symbol", "instrument", "ticker"},
    "quantity": {"quantity", "qty"},
    "settle": {"tenor", "settle", "settle date", "value date", "maturity", "maturity date"},
    "fill": {"fill", "price", "rate", "fill price", "entry price"},
}

_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_US_DATE_RE = re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")


@dataclass(frozen=True)
class RowIssue:
    """A single skipped row, in language a non-technical reader can act on."""
    row: int
    column: str
    issue: str
    raw_value: object

    @property
    def message(self) -> str:
        return (f"Row {self.row}, column '{self.column}': {self.issue}, "
                f"found {self.raw_value!r} — this row was skipped.")


@dataclass(frozen=True)
class FuturesFill:
    row: int
    trade_date: str
    pair: str          # e.g. 'ESU6 Index'
    root: str           # e.g. 'ES'
    contracts: float
    fill: float
    settle_date: str    # value/expiry date saved in the workbook


class FillList(list):
    """A list of FuturesFill that also carries the diagnostics for skipped rows."""

    def __init__(self, fills, issues):
        super().__init__(fills)
        self.issues: List[RowIssue] = list(issues)

    @property
    def messages(self) -> List[str]:
        return [i.message for i in self.issues]


def _normalize(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def _find_sheet(book, name: str):
    target = _normalize(name)
    for title in book.sheetnames:
        if _normalize(title) == target:
            return book[title]
    raise ValueError(
        f"Workbook has no sheet named '{name}' (case/whitespace-insensitive match); "
        f"found sheets: {', '.join(book.sheetnames)}")


def _find_header_row(sheet, max_scan: int = 5) -> int:
    """Blank leading rows are tolerated: scan the first few rows for a 'Date' header."""
    for r in range(1, min(max_scan, sheet.max_row) + 1):
        for c in range(1, sheet.max_column + 1):
            if _normalize(sheet.cell(r, c).value) in _HEADER_ALIASES["date"]:
                return r
    return 1


def _column_map(sheet, header_row: int) -> Dict[str, int]:
    headers = {}
    for col in range(1, sheet.max_column + 1):
        norm = _normalize(sheet.cell(header_row, col).value)
        if norm and norm not in headers:  # first match wins on duplicate headers
            headers[norm] = col
    colmap = {}
    for field_name, aliases in _HEADER_ALIASES.items():
        found = next((headers[a] for a in aliases if a in headers), None)
        colmap[field_name] = found if found is not None else _DEFAULT_COLUMNS[field_name]
    return colmap


def _try_iso(value) -> Tuple[Optional[str], Optional[str]]:
    """Best-effort date normalisation. Returns (iso_string, None) or (None, plain-English error)."""
    if isinstance(value, datetime):
        return value.date().isoformat(), None
    if isinstance(value, date):
        return value.isoformat(), None
    if isinstance(value, bool):
        return None, "expected a date, found a boolean"
    if isinstance(value, (int, float)):
        try:
            return xl_datetime.from_excel(value).date().isoformat(), None
        except Exception:
            return None, f"expected a date, found the number {value!r} which is not a valid Excel date serial"
    text = str(value).strip()
    if _ISO_DATE_RE.match(text):
        try:
            return date.fromisoformat(text).isoformat(), None
        except ValueError:
            return None, "expected a valid date"
    if _US_DATE_RE.match(text):
        try:
            return datetime.strptime(text, "%m/%d/%Y").date().isoformat(), None
        except ValueError:
            return None, "expected a valid date"
    return None, "expected a date"


def _as_iso(value) -> str:
    """Strict variant used once a value is already known-good (post row-issue checks)."""
    iso, err = _try_iso(value)
    if err:
        raise ValueError(err)
    return iso


def _extract_contracts(cell_value, fill: float) -> float:
    if hasattr(cell_value, "text"):  # openpyxl ArrayFormula
        cell_value = cell_value.text
    if isinstance(cell_value, str) and cell_value.strip().startswith("="):
        m = FUTURES_FORMULA_RE.match(cell_value.replace(" ", ""))
        if not m:
            raise ValueError(f"unrecognised futures Quantity formula {cell_value!r}")
        return float(m.group(1) or m.group(2))
    # No formula saved: only the computed value is on disk. Recover contracts from it.
    if fill == 0:
        raise ValueError("fill is zero, cannot recover contracts from a bare value")
    try:
        return float(cell_value) / (50.0 * fill)
    except (TypeError, ValueError):
        raise ValueError("expected a number or a futures Quantity formula")


def _expiry_for(pair: str, trade_date: date, row: int) -> str:
    m = INDEX_PAIR_RE.match(pair)
    if not m:
        raise ValueError(f"row {row}: unrecognised futures pair {pair!r}, expected '<ROOT><code><digit> Index'")
    root, mcode, ydigit = m.groups()
    month = MONTH_CODES.index(mcode) + 1
    year = (trade_date.year // 10) * 10 + int(ydigit)
    if year < trade_date.year - 1:
        year += 10
    return third_friday(year, month).isoformat()


def read_futures_fills(xlsx_path: Union[str, Path]) -> FillList:
    """Pure parse: every row of 'All FX trades' whose pair column ends in ' Index'.

    Cosmetic format variation (sheet name case, header order/whitespace, blank padding
    rows, date representation) is tolerated. Genuinely bad data is never guessed at: the
    row is skipped and recorded in the returned FillList's ``.issues``.
    """
    book = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    try:
        sheet = _find_sheet(book, SHEET)
        header_row = _find_header_row(sheet)
        cols = _column_map(sheet, header_row)
        out: List[FuturesFill] = []
        issues: List[RowIssue] = []
        for row_no in range(header_row + 1, sheet.max_row + 1):
            pair = sheet.cell(row_no, cols["pair"]).value
            if isinstance(pair, str):
                pair = pair.strip()
            if not isinstance(pair, str) or not pair:
                continue  # blank padding row
            if not pair.endswith(" Index"):
                continue  # an FX row of this sheet, deliberately not imported here

            trade_date_raw = sheet.cell(row_no, cols["date"]).value
            fill_raw = sheet.cell(row_no, cols["fill"]).value
            settle_raw = sheet.cell(row_no, cols["settle"]).value
            qty_raw = sheet.cell(row_no, cols["quantity"]).value

            if trade_date_raw is None or fill_raw is None or settle_raw is None or qty_raw is None:
                issues.append(RowIssue(row_no, "row", "expected Date/Quantity/tenor/fill all present",
                                        (trade_date_raw, qty_raw, settle_raw, fill_raw)))
                continue

            trade_date_iso, err = _try_iso(trade_date_raw)
            if err:
                issues.append(RowIssue(row_no, "Date", err, trade_date_raw))
                continue

            settle_iso, err = _try_iso(settle_raw)
            if err:
                issues.append(RowIssue(row_no, "tenor", err, settle_raw))
                continue

            try:
                fill = float(fill_raw)
            except (TypeError, ValueError):
                issues.append(RowIssue(row_no, "fill", "expected a number", fill_raw))
                continue

            try:
                contracts = _extract_contracts(qty_raw, fill)
            except ValueError as exc:
                issues.append(RowIssue(row_no, "Quantity", str(exc), qty_raw))
                continue

            m = INDEX_PAIR_RE.match(pair)
            if not m:
                issues.append(RowIssue(
                    row_no, "pair",
                    "expected a futures symbol shaped '<ROOT><month code><year digit> Index'", pair))
                continue
            root = m.group(1)

            out.append(FuturesFill(
                row=row_no, trade_date=trade_date_iso, pair=pair, root=root,
                contracts=contracts, fill=fill, settle_date=settle_iso,
            ))
        return FillList(out, issues)
    finally:
        book.close()


class ImportResult(int):
    """Behaves like the plain inserted-row count callers already format into messages,
    while also exposing the row-level diagnostics for a richer upload-summary panel."""

    def __new__(cls, inserted: int, issues: List[RowIssue]):
        obj = int.__new__(cls, inserted)
        obj.issues = list(issues)
        return obj

    @property
    def messages(self) -> List[str]:
        return [i.message for i in self.issues]


def load_futures_fills(xlsx_path: Union[str, Path], conn: sqlite3.Connection) -> ImportResult:
    """Insert xlsx futures fills as trades/trade_legs. Idempotent: trade_id is
    deterministic ('XL-<row>'), so a row already present is left untouched.

    Rows that fail to parse are skipped (see read_futures_fills) rather than aborting the
    whole import; the diagnostics are available on the returned ImportResult.issues."""
    schema.create_schema(conn)
    fills = read_futures_fills(xlsx_path)
    existing = {r[0] for r in conn.execute("SELECT trade_id FROM trades WHERE source = 'XLSX'")}
    inherited = dict(conn.execute("SELECT instrument_id, theme FROM instrument_theme"))
    inserted = 0
    with conn:
        for f in fills:
            trade_id = f"XL-{f.row}"
            if trade_id in existing:
                continue
            instrument_id = f.pair
            multiplier = 50.0
            conn.execute(
                "INSERT OR IGNORE INTO instruments VALUES (?,?,?,?,?,?,?,?)",
                (instrument_id, "FUTURE", f.root, "USD", multiplier, 0, instrument_id,
                 _expiry_for(f.pair, date.fromisoformat(f.trade_date), f.row)))
            conn.execute(
                "INSERT INTO trades VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (trade_id, "XLSX", instrument_id, "FUTURE", trade_id, f.trade_date,
                 f.contracts, f.fill, "", "", "", "", f"XLSX futures fill row {f.row}",
                 inherited.get(instrument_id, "")))
            conn.execute(
                "INSERT INTO trade_legs VALUES (?,?,?,?,?,?,?,?,?)",
                (trade_id, 1, "NOTIONAL", "USD", f.contracts * multiplier * f.fill,
                 f.trade_date, f.settle_date, 0.0, 0))
            inserted += 1
    return ImportResult(inserted, fills.issues)
